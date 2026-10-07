import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import export_stats as ex  # noqa: E402

REPO = HERE.parents[2]
WORKBOOK = REPO / "data" / "source" / "ASL_Complete_S1_S21.xlsx"


class ParserTests(unittest.TestCase):
    def test_placement_ranges_normalise_to_en_dash(self):
        for raw in ("9th-12th", "9th–12th", "9th — 12th"):
            self.assertEqual(
                ex.to_placement(raw),
                {"label": "9th–12th", "best": 9, "worst": 12, "status": "final"},
            )

    def test_placement_singles_and_ordinals(self):
        self.assertEqual(ex.to_placement("1st")["label"], "1st")
        self.assertEqual(ex.to_placement("3rd")["best"], 3)
        self.assertEqual(ex.to_placement("23rd-28th")["label"], "23rd–28th")
        self.assertEqual(ex.to_placement("11th-12th")["label"], "11th–12th")

    def test_placement_blank_and_in_progress(self):
        for raw in (None, "-", "", "  "):
            self.assertIsNone(ex.to_placement(raw))
        self.assertEqual(ex.to_placement("In Prog")["status"], "in_progress")
        self.assertEqual(ex.to_placement("Ro16 (In Progress)")["status"], "in_progress")

    def test_placement_rejects_garbage(self):
        with self.assertRaises(ValueError):
            ex.to_placement("Ro16")

    def test_ratio(self):
        self.assertAlmostEqual(ex.to_ratio("40.0%"), 0.4)
        self.assertAlmostEqual(ex.to_ratio("0%"), 0.0)
        self.assertAlmostEqual(ex.to_ratio(0.25), 0.25)
        self.assertAlmostEqual(ex.to_ratio(45), 0.45)
        with self.assertRaises(ValueError):
            ex.to_ratio("40.0")

    def test_record(self):
        self.assertEqual(ex.to_record("44% (4-5)"), {"wins": 4, "losses": 5, "winRate": 0.444444})
        self.assertIsNone(ex.to_record("--"))
        self.assertIsNone(ex.to_record(None))
        with self.assertRaises(ValueError):
            ex.to_record("44%")

    def test_ints_reject_fractions_and_strip_currency(self):
        self.assertEqual(ex.to_int(3.0), 3)
        self.assertEqual(ex.to_int("₩1,500,000"), 1_500_000)
        self.assertIsNone(ex.to_int("TBD"))
        with self.assertRaises(ValueError):
            ex.to_int(2.5)

    def test_race_codes(self):
        self.assertEqual(ex.to_race("Terran"), "T")
        self.assertEqual(ex.to_race("Z"), "Z")
        self.assertIsNone(ex.to_race("?"))
        with self.assertRaises(ValueError):
            ex.to_race("Random")

    def test_season_headers(self):
        self.assertEqual(ex.to_season_number("S18\n(SSL)"), 18)
        self.assertEqual(ex.to_season_number("S21\nProg"), 21)

    def test_names_treat_placeholders_as_unknown(self):
        self.assertIsNone(ex.to_name("TBD"))
        self.assertEqual(ex.to_name(" Flash "), "Flash")


class ValidationTests(unittest.TestCase):
    def _data(self, **over):
        base = {
            "seasons": [{"season": 1, "winner": "Flash", "runnerUp": "Sea"}],
            "placements": [
                {"player": "Flash", "race": "T", "season": 1, "placement": ex.to_placement("1st")},
                {"player": "Sea", "race": "T", "season": 1, "placement": ex.to_placement("2nd")},
            ],
            "elo": [{"player": "Flash", "race": "T"}, {"player": "Sea", "race": "T"}],
            "playerStats": [
                {"player": "Flash", "race": "T", "championships": 1, "estPrizeKrw": 10_000_000},
                {"player": "Sea", "race": "T", "championships": 0, "estPrizeKrw": 5_000_000},
            ],
            "liveTracker": [],
        }
        base.update(over)
        return base

    def codes(self, data):
        return [w["code"] for w in ex.validate(data)]

    def test_clean_data_has_no_warnings(self):
        self.assertEqual(self.codes(self._data()), [])

    def test_detects_case_variant_identities(self):
        d = self._data()
        d["elo"].append({"player": "flash", "race": "T"})
        self.assertIn("PLAYER_CASE_VARIANTS", self.codes(d))

    def test_detects_race_conflict(self):
        d = self._data(liveTracker=[{"player": "Flash", "race": "Z", "status": "Ro16 (In Progress)"}])
        self.assertIn("RACE_CONFLICT", self.codes(d))

    def test_detects_finals_and_championship_mismatch(self):
        d = self._data(seasons=[{"season": 1, "winner": "Sea", "runnerUp": "Flash"}])
        codes = self.codes(d)
        self.assertIn("FINALS_MISMATCH", codes)
        self.assertIn("CHAMPIONSHIP_MISMATCH", codes)

    def test_detects_unscaled_prize(self):
        d = self._data()
        d["playerStats"][1]["estPrizeKrw"] = 4
        self.assertIn("SUSPECT_VALUE", self.codes(d))


@unittest.skipUnless(WORKBOOK.exists(), "source workbook not present")
class WorkbookIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_exports_every_file_with_expected_shape(self):
        result = ex.export(WORKBOOK, self.tmp / "out")
        out = self.tmp / "out"
        for name in (
            "manifest.json", "seasons.json", "players.json", "placements.json", "elo.json",
            "player-stats.json", "race-stats.json", "live-tracker.json", "validation.json",
        ):
            self.assertTrue((out / name).exists(), name)
        seasons = json.loads((out / "seasons.json").read_text(encoding="utf-8"))
        self.assertEqual([s["season"] for s in seasons], list(range(1, len(seasons) + 1)))
        self.assertTrue(all(s["status"] in ("complete", "in_progress") for s in seasons))
        placements = json.loads((out / "placements.json").read_text(encoding="utf-8"))
        self.assertTrue(all("–" in p["placement"]["label"] or p["placement"]["best"] == p["placement"]["worst"]
                            or p["placement"]["status"] == "in_progress" for p in placements))
        self.assertEqual(result["manifest"]["counts"]["players"], len(json.loads((out / "players.json").read_text(encoding="utf-8"))))

    def test_export_is_deterministic_across_processes(self):
        # Separate interpreters with different hash seeds: set/dict ordering must not leak into output.
        script = HERE.parent / "export_stats.py"
        for seed, name in (("1", "a"), ("2", "b")):
            env = {**os.environ, "PYTHONHASHSEED": seed}
            subprocess.run([sys.executable, str(script), str(WORKBOOK), str(self.tmp / name)],
                           check=True, capture_output=True, env=env)
        for f in (self.tmp / "a").iterdir():
            self.assertEqual(f.read_bytes(), (self.tmp / "b" / f.name).read_bytes(), f.name)

    def test_renamed_header_fails_loudly(self):
        broken = self.tmp / "broken.xlsx"
        wb = openpyxl.load_workbook(WORKBOOK)
        wb["Race Stats"]["E4"] = "Win %"
        wb.save(broken)
        with self.assertRaises(ex.LayoutError) as ctx:
            ex.export(broken, self.tmp / "out")
        self.assertIn("overall", str(ctx.exception))

    def test_missing_sheet_fails_loudly(self):
        broken = self.tmp / "broken.xlsx"
        wb = openpyxl.load_workbook(WORKBOOK)
        del wb["ELO Ratings"]
        wb.save(broken)
        with self.assertRaises(ex.LayoutError):
            ex.export(broken, self.tmp / "out")

    def test_cli_strict_exits_nonzero_when_warnings(self):
        code = ex.main([str(WORKBOOK), str(self.tmp / "out"), "--strict"])
        warnings = json.loads((self.tmp / "out" / "validation.json").read_text(encoding="utf-8"))
        self.assertEqual(code, 2 if warnings else 0)


if __name__ == "__main__":
    unittest.main()
