"""One-off workbook correction, 2026-10-07. Kept for the audit trail; do not re-run.

Applies Derek's decisions from Docs/reports/2026-10-07-stats-validation.md:

  * identity merges   Snow->SnOw, hero->herO, BeSt->Best, Effort/effOrt->EffOrt,
                      Hyun->HyuN, huro / Yoon Soo-chul -> tulbo
  * race corrections  sSak, Ample, Speed -> T; Shine -> Z; tulbo -> P; Jaedong -> Z (tracker)
  * EffOrt S5         9th-12th (the 23rd-28th entry is dropped)
  * S21 final results (source: Liquipedia "ASL Season 21", pasted by Derek)

Derived tabs:
  * Player Placements Played/Best and Player Stats counts are recomputed. Both were verified
    to reproduce exactly from placements before any edit.
  * ELO Ratings is recomputed for every player (pass --elo). The workbook's original
    algorithm could not be reproduced exactly; see compute_elo for the one used here.
  * Race Stats aggregate tables are NOT touched. They don't reproduce from placements and
    need a proper recompute (ROADMAP: "Compute derived stats in the pipeline").

Usage: python apply_2026_10_07_decisions.py <workbook.xlsx> [--elo]
"""

from __future__ import annotations

import copy
import itertools
import re
import sys
from collections import Counter, defaultdict

import openpyxl

RENAMES = {
    "Snow": "SnOw",
    "hero": "herO",
    "BeSt": "Best",
    "Effort": "EffOrt",
    "effOrt": "EffOrt",
    "Hyun": "HyuN",
    "huro": "tulbo",
    "Yoon Soo-chul": "tulbo",
}
RACE_FIX = {"sSak": "T", "Ample": "T", "Speed": "T", "Shine": "Z", "tulbo": "P"}
RACE_NAME = {"T": "Terran", "Z": "Zerg", "P": "Protoss"}

# S21 final placement per player (workbook style: hyphen ranges; both SF losers are "3rd").
S21 = {
    "soma": "1st", "Flash": "2nd", "Leta": "3rd", "Light": "3rd",
    "herO": "5th-8th", "tulbo": "5th-8th", "Jaedong": "5th-8th", "SnOw": "5th-8th",
    "sSak": "9th-12th", "Rain": "9th-12th", "Bisu": "9th-12th", "RoyaL": "9th-12th",
    "Sharp": "13th-16th", "PianO": "13th-16th", "Ample": "13th-16th", "BarrackS": "13th-16th",
    "Scan": "17th-22nd", "Soulkey": "17th-22nd", "Shine": "17th-22nd",
    "Mind": "17th-22nd", "Speed": "17th-22nd", "ZerO": "17th-22nd",
    "Mong": "23rd-28th", "JyJ": "23rd-28th", "Larva": "23rd-28th",
    "Calm": "23rd-28th", "Rush": "23rd-28th", "Best": "23rd-28th",
}
# Tracker rows in final order: (player, prize, note). Prizes for 1st-4th are from Liquipedia.
# Ro16/Ro24 prizes are the workbook's own. 5th-8th is inferred: it is the only value that
# makes the 28 payouts sum to the W78,000,000 pool.
TRACKER = [
    ("soma", "₩30,000,000", "Champion. Beat Flash 4–3 in the final; back-to-back titles (S20, S21)"),
    ("Flash", "₩10,000,000", "Runner-up. Lost the final 3–4 to soma"),
    ("Leta", "₩3,000,000", "Lost semifinal 1–4 to soma"),
    ("Light", "₩3,000,000", "Lost semifinal 0–4 to Flash"),
    ("herO", "₩2,000,000", "Lost quarterfinal 0–3 to soma"),
    ("tulbo", "₩2,000,000", "Lost quarterfinal 0–3 to Leta"),
    ("Jaedong", "₩2,000,000", "Lost quarterfinal 2–3 to Light"),
    ("SnOw", "₩2,000,000", "Lost quarterfinal 2–3 to Flash"),
    ("sSak", "₩1,500,000", "Eliminated Ro16 Group A (3rd)"),
    ("Rain", "₩1,500,000", "Eliminated Ro16 Group B (3rd)"),
    ("Bisu", "₩1,500,000", "Eliminated Ro16 Group C (3rd)"),
    ("RoyaL", "₩1,500,000", "Eliminated Ro16 Group D (3rd)"),
    ("Sharp", "₩1,500,000", "Eliminated Ro16 Group A (4th)"),
    ("PianO", "₩1,500,000", "Eliminated Ro16 Group B (4th)"),
    ("Ample", "₩1,500,000", "Eliminated Ro16 Group C (4th)"),
    ("BarrackS", "₩1,500,000", "Eliminated Ro16 Group D (4th)"),
    ("Scan", "₩1,000,000", "Eliminated Ro24 Group A (3rd)"),
    ("Soulkey", "₩1,000,000", "Eliminated Ro24 Group B (3rd); three-time defending finalist out early"),
    ("Shine", "₩1,000,000", "Eliminated Ro24 Group C (3rd)"),
    ("Mind", "₩1,000,000", "Eliminated Ro24 Group D (3rd)"),
    ("Speed", "₩1,000,000", "Eliminated Ro24 Group E (3rd)"),
    ("ZerO", "₩1,000,000", "Eliminated Ro24 Group F (3rd)"),
    ("Mong", "₩1,000,000", "Eliminated Ro24 Group A (4th)"),
    ("JyJ", "₩1,000,000", "Eliminated Ro24 Group B (4th)"),
    ("Larva", "₩1,000,000", "Eliminated Ro24 Group C (4th)"),
    ("Calm", "₩1,000,000", "Eliminated Ro24 Group D (4th)"),
    ("Rush", "₩1,000,000", "Eliminated Ro24 Group E (4th)"),
    ("Best", "₩1,000,000", "Eliminated Ro24 Group F (4th)"),
]

FIRST_SEASON_COL, S21_COL, PLAYED_COL, BEST_COL = 3, 23, 24, 25  # C, W, X, Y
K = 32


def tier(v) -> int | None:
    if v in (None, "-"):
        return None
    m = re.match(r"(\d+)", str(v))
    if not m:
        raise ValueError(f"not a final placement: {v!r}")
    return int(m.group(1))


def restyle(dst, src) -> None:
    dst._style = copy.copy(src._style)


def compute_elo(tiers: dict[str, list[int | None]]):
    """Placement-based ELO, start 1500, K=32, S1-S20.

    Each season, participants are ordered by (placement, name). Every pair in different
    placement tiers is one game won by the better-placed player; pairs in the same tier are
    skipped. Ratings update after each pair. Peak is the best rating after any season played.
    """
    rating = {n: 1500.0 for n in tiers}
    peak: dict[str, tuple[float, int]] = {}
    for s in range(20):
        part = sorted((n for n in tiers if tiers[n][s] is not None), key=lambda n: (tiers[n][s], n.lower(), n))
        for a, b in itertools.combinations(part, 2):
            if tiers[a][s] == tiers[b][s]:
                continue
            delta = K * (1 - 1 / (1 + 10 ** ((rating[b] - rating[a]) / 400)))
            rating[a] += delta
            rating[b] -= delta
        for n in part:
            if n not in peak or rating[n] > peak[n][0]:
                peak[n] = (rating[n], s + 1)
    return rating, peak


def main(path: str, do_elo: bool) -> None:
    wb = openpyxl.load_workbook(path)
    so, pp, el, ps, rs, lt = (wb[n] for n in (
        "Season Overview", "Player Placements", "ELO Ratings", "Player Stats", "Race Stats", "S21 Live Tracker"))
    if not any(pp.cell(r, 1).value == "BeSt" for r in range(3, pp.max_row + 1)):
        sys.exit("already applied (no 'BeSt' row); refusing to run twice")

    # ---- capture what we need from the untouched workbook
    old_prize = {ps.cell(r, 1).value: ps.cell(r, 9).value for r in range(3, ps.max_row + 1) if ps.cell(r, 1).value}
    old_elo = {el.cell(r, 2).value: el.cell(r, 4).value for r in range(3, el.max_row + 1) if el.cell(r, 2).value}
    cell_style: dict[tuple, Counter] = defaultdict(Counter)  # (value, race) and (value, None) -> style
    style_obj = {}
    best_style: dict[str, object] = {}
    name_style: dict[str, tuple] = {}
    for r in range(3, pp.max_row + 1):
        name, race = pp.cell(r, 1).value, pp.cell(r, 2).value
        if not name or name in RACE_FIX or RENAMES.get(name) in RACE_FIX:
            continue
        name_style.setdefault(race, (pp.cell(r, 1), pp.cell(r, 2)))
        for c in range(FIRST_SEASON_COL, S21_COL + 1):
            cell = pp.cell(r, c)
            if cell.value in (None, "-", "In Prog"):
                continue
            key = tuple(cell._style)
            style_obj[key] = cell
            cell_style[(cell.value, race)][key] += 1
            cell_style[(cell.value, None)][key] += 1
        best_style.setdefault(pp.cell(r, BEST_COL).value, pp.cell(r, BEST_COL))

    def placement_style(value, race):
        for k in ((value, race), (value, None)):
            if cell_style.get(k):
                return style_obj[cell_style[k].most_common(1)[0][0]]
        return None

    # ---- 1. renames, every sheet, exact cell match
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in RENAMES:
                    cell.value = RENAMES[cell.value]

    # ---- 2. Player Placements: merge duplicate rows
    rows_by_name: dict[str, list[int]] = defaultdict(list)
    for r in range(3, pp.max_row + 1):
        if pp.cell(r, 1).value:
            rows_by_name[pp.cell(r, 1).value].append(r)
    to_delete: list[int] = []
    for name, rows in rows_by_name.items():
        if len(rows) == 1:
            continue
        target = rows[0]
        for c in range(FIRST_SEASON_COL, S21_COL + 1):
            vals = [(r, pp.cell(r, c).value) for r in rows if pp.cell(r, c).value not in (None, "-")]
            if len(vals) > 1:
                if (name, c) != ("EffOrt", FIRST_SEASON_COL + 4):  # S5
                    sys.exit(f"unexpected double placement for {name} in column {c}: {vals}")
                vals = [(r, v) for r, v in vals if v == "9th-12th"]
            if vals and vals[0][0] != target:
                src = pp.cell(vals[0][0], c)
                pp.cell(target, c).value = src.value
                restyle(pp.cell(target, c), src)
        to_delete += rows[1:]
    for r in sorted(to_delete, reverse=True):
        pp.delete_rows(r)

    # ---- 3. race corrections + S21 results + Played/Best
    race_of: dict[str, str] = {}
    tiers: dict[str, list[int | None]] = {}
    for r in range(3, pp.max_row + 1):
        name = pp.cell(r, 1).value
        if not name:
            continue
        if name in RACE_FIX:
            race = RACE_FIX[name]
            pp.cell(r, 2).value = race
            restyle(pp.cell(r, 1), name_style[race][0])
            restyle(pp.cell(r, 2), name_style[race][1])
            for c in range(FIRST_SEASON_COL, S21_COL + 1):
                if pp.cell(r, c).value not in (None, "-"):
                    src = placement_style(pp.cell(r, c).value, race)
                    if src is not None:
                        restyle(pp.cell(r, c), src)
        race = race_of[name] = pp.cell(r, 2).value
        w = pp.cell(r, S21_COL)
        want = S21.get(name, "-")
        if w.value != want:
            if want == "-":
                sys.exit(f"{name} has an S21 entry {w.value!r} but is not in the final results")
            w.value = want
            src = placement_style(want, race)
            if src is not None:
                restyle(w, src)
        vals = [pp.cell(r, c).value for c in range(FIRST_SEASON_COL, S21_COL + 1)]
        played = [v for v in vals if v not in (None, "-")]
        best = min(played, key=tier)
        pp.cell(r, PLAYED_COL).value = len(played)
        if pp.cell(r, BEST_COL).value != best:
            pp.cell(r, BEST_COL).value = best
            if best in best_style:
                restyle(pp.cell(r, BEST_COL), best_style[best])
        tiers[name] = [tier(v) for v in vals[:20]]
    missing = set(S21) - set(race_of)
    if missing:
        sys.exit(f"S21 players without a Player Placements row: {sorted(missing)}")

    # ---- 4. Season Overview
    for cell, ref in ((so["I5"], so["I6"]), (so["J5"], so["J6"])):  # S3 runner-up Shine is Zerg
        restyle(cell, ref)
    so["J5"].value = "Z"
    so["E23"].value, so["F23"].value = 78_000_000, 51_303
    so["G23"].value, so["H23"].value = "soma", "Z"
    so["I23"].value, so["J23"].value = "Flash", "T"
    so["K23"].value = "soma back-to-back (S20, S21)"
    for col, ref_row in (("G", 22), ("H", 22), ("I", 8), ("J", 8)):  # Zerg winner, Terran runner-up
        restyle(so[f"{col}23"], so[f"{col}{ref_row}"])

    # ---- 5. Race Stats: champion reference table only (names were renamed above)
    rs["Q13"].value = "Zerg"
    restyle(rs["P13"], rs["P14"])
    restyle(rs["Q13"], rs["Q14"])

    # ---- 6. S21 tracker: final standings for all 28 players
    row_of = {lt.cell(r, 1).value: r for r in range(3, 29)}
    # Eliminated-row look: race column is race-coloured, the rest is zebra-striped by row.
    race_tmpl = {race: copy.copy(lt.cell(row_of[n], 2)._style) for race, n in (("T", "Sharp"), ("Z", "Soulkey"), ("P", "Rain"))}
    stripe = {1: [copy.copy(lt.cell(11, c)._style) for c in range(1, 6)],
              0: [copy.copy(lt.cell(12, c)._style) for c in range(1, 6)]}
    note_style, blank_style = copy.copy(lt["A30"]._style), [copy.copy(lt.cell(29, c)._style) for c in range(1, 6)]
    lt.unmerge_cells("A30:E30")
    lt["A1"].value = "ASL Season 21 Tracker (Final Results)"
    for i, (name, prize, note) in enumerate(TRACKER):
        r, race = 3 + i, race_of[name]
        label = S21[name].replace("-", "–")
        for c, v in enumerate((name, race, label, prize, note), start=1):
            lt.cell(r, c).value = v
            lt.cell(r, c)._style = copy.copy(race_tmpl[race] if c == 2 else stripe[r % 2][c - 1])
    for c in range(1, 6):
        lt.cell(31, c).value = None
        lt.cell(31, c)._style = copy.copy(blank_style[c - 1])
        lt.cell(32, c)._style = copy.copy(note_style)
    lt["A32"].value = ("Last Updated: October 7, 2026. Season complete (final played May 24, 2026). "
                       "Results from Liquipedia; the 5th–8th prize is inferred from the total pool.")
    lt.merge_cells("A32:E32")

    # ---- 7. ELO Ratings (optional) and Player Stats
    if do_elo:
        rating, peak = compute_elo(tiers)
        order = sorted((n for n in tiers if n in peak), key=lambda n: (-round(rating[n], 1), n.lower(), n))
        name_t = {race_of_old: st for race_of_old, st in ((el.cell(r, 3).value, copy.copy(el.cell(r, 2)._style)) for r in range(el.max_row, 5, -1))}
        race_t = {el.cell(r, 3).value: copy.copy(el.cell(r, 3)._style) for r in range(el.max_row, 5, -1)}
        champ_t = {bool(el.cell(r, 8).value): copy.copy(el.cell(r, 8)._style) for r in range(el.max_row, 5, -1)}
        old_rows = sum(1 for r in range(3, el.max_row + 1) if el.cell(r, 2).value)
        for i, n in enumerate(order):
            r = 3 + i
            champs = sum(1 for t in tiers[n] if t == 1)
            vals = (i + 1, n, race_of[n], round(rating[n], 1), round(peak[n][0], 1), f"S{peak[n][1]}",
                    sum(1 for t in tiers[n] if t is not None), champs)
            for c, v in enumerate(vals, start=1):
                el.cell(r, c).value = v
            el.cell(r, 2)._style = copy.copy(name_t[race_of[n]])
            el.cell(r, 3)._style = copy.copy(race_t[race_of[n]])
            el.cell(r, 8)._style = copy.copy(champ_t.get(bool(champs), champ_t[False]))
            if r <= 5:
                f = copy.copy(el.cell(r, 2).font)
                f.b = True
                el.cell(r, 2).font = f
        if old_rows > len(order):
            el.delete_rows(3 + len(order), old_rows - len(order))
        new_elo = {n: round(rating[n], 1) for n in order}
        same = [n for n in new_elo if n in old_elo and n not in RENAMES.values()]
        print("ELO recomputed: %d players; unmerged players moved by median %.1f, max %.1f (%s)" % (
            len(order),
            sorted(abs(new_elo[n] - old_elo[n]) for n in same)[len(same) // 2],
            max(abs(new_elo[n] - old_elo[n]) for n in same),
            max(same, key=lambda n: abs(new_elo[n] - old_elo[n]))))
    else:
        new_elo = {el.cell(r, 2).value: el.cell(r, 4).value for r in range(3, el.max_row + 1) if el.cell(r, 2).value}

    # Player Stats: counts from placements (S1-S20); prize carried over, summed for merges.
    prize = {RENAMES.get(n, n): 0 for n in old_prize}
    for n, v in old_prize.items():
        prize[RENAMES.get(n, n)] += v or 0
    prize["Best"] = None  # the old 'Best' row held 4 (lost scale); Derek to supply the real total
    name_t = {ps.cell(r, 2).value: copy.copy(ps.cell(r, 1)._style) for r in range(ps.max_row, 5, -1)}
    race_t = {ps.cell(r, 2).value: copy.copy(ps.cell(r, 2)._style) for r in range(ps.max_row, 5, -1)}
    old_rows = sum(1 for r in range(3, ps.max_row + 1) if ps.cell(r, 1).value)
    order = sorted((n for n in tiers if any(t is not None for t in tiers[n]) and n in new_elo),
                   key=lambda n: (-new_elo[n], n.lower(), n))
    labels = {}
    for r in range(3, pp.max_row + 1):
        n = pp.cell(r, 1).value
        if n:
            vals = [pp.cell(r, c).value for c in range(FIRST_SEASON_COL, S21_COL)]
            labels[n] = [v for v in vals if v not in (None, "-")]
    for i, n in enumerate(order):
        r, t = 3 + i, [x for x in tiers[n] if x is not None]
        vals = (n, race_of[n], len(t), min(labels[n], key=tier), t.count(1), sum(x <= 2 for x in t),
                sum(x <= 4 for x in t), sum(x <= 5 for x in t), prize.get(n), new_elo[n])
        for c, v in enumerate(vals, start=1):
            ps.cell(r, c).value = v
        ps.cell(r, 1)._style = copy.copy(name_t[race_of[n]])
        ps.cell(r, 2)._style = copy.copy(race_t[race_of[n]])
    if old_rows > len(order):
        ps.delete_rows(3 + len(order), old_rows - len(order))

    wb.save(path)
    per_season = [sum(1 for n in tiers if tiers[n][s] is not None) for s in range(20)]
    print("players:", len(race_of), "| S21 entries:", sum(1 for n in race_of if n in S21),
          "| seasons with != 28 players (S1 has 16):", {f"S{s + 1}": c for s, c in enumerate(per_season) if c != 28})


if __name__ == "__main__":
    main(sys.argv[1], "--elo" in sys.argv[2:])
